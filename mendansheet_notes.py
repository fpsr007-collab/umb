from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def heading1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '2E74B5')
    pPr.append(shd)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.size = Pt(10.5)
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.size = Pt(10.5)
    else:
        r = p.add_run(text)
        r.font.size = Pt(10.5)
    return p

def note_box(text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    cell.text = text
    set_cell_bg(cell, 'FFF2CC')
    for run in cell.paragraphs[0].runs:
        run.font.size = Pt(10)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return tbl

# ===== タイトル =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run('仕事と介護の両立支援プラン（面談シート兼用）')
tr.bold = True
tr.font.size = Pt(15)
tr.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
title.paragraph_format.space_after = Pt(2)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = subtitle.add_run('作成・面談実施の注意点')
sr.font.size = Pt(11)
sr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
subtitle.paragraph_format.space_after = Pt(14)

# ===== 1. 面談の実施者・回数 =====
heading1('１．面談の実施者・回数')

bullet('実施者は対象労働者の上司または人事労務担当者と対象労働者の組み合わせで行う', '実施者：')
bullet('少なくとも1回以上実施し、結果を必ず面談シートに記録する', '回数：')
bullet('対面が困難な場合は電話・メール等による相談・調整の記録でも可', '対面困難時：')

# ===== 2. 作成タイミング =====
heading1('２．作成タイミング（最重要）')

tbl1 = doc.add_table(rows=3, cols=4)
tbl1.style = 'Table Grid'

h1 = ['種別', '原則', '例外（可）', '不支給になるケース']
for i, h in enumerate(h1):
    cell = tbl1.cell(0, i)
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(10)
    set_cell_bg(cell, '2E74B5')
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

timing_data = [
    ('介護休業',
     '休業開始前に作成',
     '休業期間中の作成も可',
     '休業終了後にプラン作成・面談を行った場合'),
    ('介護両立支援制度',
     '制度利用開始前日までに作成',
     '制度利用期間中の作成も可',
     '制度利用終了後にプラン作成・面談を行った場合'),
]

col_widths = [Cm(2.8), Cm(3.5), Cm(4.0), Cm(5.2)]
for i, w in enumerate(col_widths):
    for row in tbl1.rows:
        row.cells[i].width = w

for i, row_data in enumerate(timing_data):
    row = tbl1.rows[i+1]
    for j, val in enumerate(row_data):
        row.cells[j].text = val
        for run in row.cells[j].paragraphs[0].runs:
            run.font.size = Pt(10)
    if i % 2 == 0:
        for j in range(4):
            set_cell_bg(row.cells[j], 'DEEAF1')
    # 不支給列は赤系
    row.cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

note_box('⚠ プランの作成や面談が休業・制度利用の「終了後」になった場合は一切支給対象外となるため注意。')
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ===== 3. プランに記載すべき内容 =====
heading1('３．プランに記載すべき内容')

bullet('【介護休業の場合】対象労働者の業務の整理・引き継ぎに関する措置を少なくとも記載する')
bullet('【介護両立支援制度の場合】制度利用期間中の業務体制の検討に関する取組を記載する')
bullet('プランに基づく業務の整理・引き継ぎが行われないまま休業が終了した場合は支給対象外')
bullet('引き継ぎが困難な場合は電話・メール・書面による引き継ぎでも可（対面不要）')

# ===== 4. フォロー面談 =====
heading1('４．フォロー面談（介護休業のみ）')

bullet('介護休業終了後に上司または人事労務担当者とフォロー面談を実施し、結果を面談シートに記録する')
bullet('原職等に復帰させることが原則。ただし労働者本人の希望により原職等以外で復帰する場合は、希望が面談記録等で確認できることが必要')
bullet('無期雇用労働者が介護休業後に有期雇用契約を締結している場合は対象外（本人希望でも不可）')

# ===== 5. 複数制度・複数回利用時 =====
heading1('５．複数制度・複数回利用時の注意')

bullet('同一対象家族について異なる複数の介護両立支援制度を利用する場合は、制度ごとに新たに面談を実施しプランを作成する', '')
bullet('制度利用開始後に当初予定していなかった別の制度を利用することになった場合も、面談シートに利用した全制度と利用期間が確認できることが必要')
bullet('同一労働者・同一対象家族への支給は合計2回まで（3種類目の制度利用は対象外）')
bullet('介護休業取得後に同一対象家族について両立支援制度を利用する場合も、新たに面談を実施しプランを作成する必要がある')

tbl2 = doc.add_table(rows=3, cols=2)
tbl2.style = 'Table Grid'

multi_headers = ['ケース', '対応']
for i, h in enumerate(multi_headers):
    cell = tbl2.cell(0, i)
    cell.text = h
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(10)
    set_cell_bg(cell, '2E74B5')
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

multi_data = [
    ('同一対象家族について制度Aを利用後、同一制度Aを再利用',
     '対象外（同一制度は1回限り）'),
    ('同一対象家族について制度Aを利用後、別の制度Bを利用',
     'A・B それぞれ対象（合計2回まで）。制度Bについて新たに面談・プラン作成が必要'),
]

for i, (k, v) in enumerate(multi_data):
    row = tbl2.rows[i+1]
    row.cells[0].text = k
    row.cells[1].text = v
    for cell in row.cells:
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(10)
    if i % 2 == 0:
        for j in range(2):
            set_cell_bg(row.cells[j], 'DEEAF1')

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ===== 6. 休業期間変更時 =====
heading1('６．休業期間・制度利用期間が変更になった場合')

bullet('介護休業期間変更届の提出だけでは不足。変更後の期間についてプランへの記載（更新）も必要')
bullet('介護の状況により当初のプランから変更が生じた場合は、事業主・労働者双方のためにプライ自体を更新することが求められる')

# ===== 7. 全労働者への方針周知 =====
heading1('７．全労働者への方針周知（プラン前提条件）')

bullet('介護休業取得・職場復帰・制度利用を支援する方針を全労働者へ事前周知すること')
bullet('周知期限は休業開始日または制度利用開始日の前日まで（休業・制度利用と同時並行での実施も可）')
bullet('周知が終了後になった場合は支給対象外')
bullet('周知方法：労働協約・就業規則への規定、または方針を明文化した文書・社内報・イントラネット等')

# ===== 注釈 =====
doc.add_paragraph()
note = doc.add_paragraph()
nr = note.add_run('※ 本資料は令和8年4月8日改正版の支給要領に基づき作成しています。最新情報は管轄労働局または厚生労働省ウェブサイトでご確認ください。')
nr.font.size = Pt(9)
nr.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
nr.italic = True

output_path = '/home/user/umb/面談シート作成の注意点.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
