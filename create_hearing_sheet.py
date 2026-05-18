"""
就業規則作成用 ヒアリングシート Excel生成スクリプト
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# カラー定義
C_TITLE_BG   = "1F4E79"
C_TITLE_FG   = "FFFFFF"
C_SEC_COLORS = [
    ("2E75B6", "FFFFFF"),  # 青
    ("375623", "FFFFFF"),  # 緑
    ("7030A0", "FFFFFF"),  # 紫
    ("C55A11", "FFFFFF"),  # オレンジ
    ("833C00", "FFFFFF"),  # 茶
    ("1F4E79", "FFFFFF"),  # 濃紺
    ("4472C4", "FFFFFF"),  # 中青
    ("538135", "FFFFFF"),  # 草緑
]
C_ITEM_BG    = "D6E4F7"
C_ITEM_FG    = "1F4E79"
C_INPUT_BG   = "FFF2CC"
C_NOTE_BG    = "F2F2F2"
C_NOTE_FG    = "595959"
C_REQ_BG     = "FCE4D6"  # 必須項目


def thin_border(color="AAAAAA"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def cell_style(cell, value=None, bold=False, size=10, fg=None, bg=None,
               h_align="left", v_align="center", wrap=False,
               number_format=None, border=True, italic=False):
    if value is not None:
        cell.value = value
    cell.font = Font(bold=bold, size=size, color=fg or "000000", italic=italic)
    if bg:
        cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.alignment = Alignment(horizontal=h_align, vertical=v_align, wrap_text=wrap)
    if border:
        cell.border = thin_border()
    if number_format:
        cell.number_format = number_format


def scw(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width


def merge_range(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def section_header(ws, row, title, sec_idx):
    bg, fg = C_SEC_COLORS[sec_idx % len(C_SEC_COLORS)]
    ws.row_dimensions[row].height = 22
    merge_range(ws, row, 1, row, 5)
    cell_style(ws.cell(row, 1), value=title, bold=True, size=11,
               fg=fg, bg=bg, h_align="left", border=False)


def item_row(ws, row, no, label, input_val="", note="",
             required=False, height=18):
    ws.row_dimensions[row].height = height
    # No.
    cell_style(ws.cell(row, 1), value=no, size=9,
               h_align="center", bg="F0F0F0")
    # 項目名
    bg = C_REQ_BG if required else C_ITEM_BG
    cell_style(ws.cell(row, 2), value=label, bold=required, size=10,
               fg=C_ITEM_FG, bg=bg, wrap=True)
    # 回答欄
    cell_style(ws.cell(row, 3), value=input_val, bg=C_INPUT_BG, wrap=True)
    # 選択肢欄（col4）
    cell_style(ws.cell(row, 4), bg=C_INPUT_BG, wrap=True)
    # 記入例・備考
    cell_style(ws.cell(row, 5), value=note, size=9, fg=C_NOTE_FG,
               bg=C_NOTE_BG, wrap=True, italic=True)


def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    for c in cells:
        dv.add(c)


# ─────────────────────────────────────────────
# ヒアリングシート本体
# 列構成:
#  A(4)=No.  B(28)=項目  C(22)=回答欄  D(18)=選択肢  E(35)=記入例・備考
# ─────────────────────────────────────────────
def build_hearing_sheet(ws):
    ws.sheet_view.showGridLines = False
    scw(ws, 1, 4)
    scw(ws, 2, 28)
    scw(ws, 3, 22)
    scw(ws, 4, 18)
    scw(ws, 5, 35)

    # タイトル
    ws.row_dimensions[1].height = 32
    merge_range(ws, 1, 1, 1, 5)
    cell_style(ws.cell(1, 1),
               value="就業規則作成　ヒアリングシート",
               bold=True, size=15, fg=C_TITLE_FG, bg=C_TITLE_BG,
               h_align="center", border=False)

    ws.row_dimensions[2].height = 18
    merge_range(ws, 2, 1, 2, 5)
    cell_style(ws.cell(2, 1),
               value="※ 黄色セルにご記入ください　　■は必須項目　　不明な場合は「未定」「要検討」とご記入ください",
               size=9, fg="9C0006", bg="FCE4D6", h_align="left", border=False)

    ws.row_dimensions[3].height = 18
    merge_range(ws, 3, 1, 3, 5)
    cell_style(ws.cell(3, 1),
               value="【担当社労士記入欄】　事業所名：　　　　　　　　　　ヒアリング日：　　　　　　　担当：",
               size=10, bg="EBF3FB", border=False)

    ws.row_dimensions[4].height = 5  # スペース

    # ── 列ヘッダー ──
    ws.row_dimensions[5].height = 18
    for col, label in [(1, "No."), (2, "項　目"), (3, "回　答"),
                       (4, "選択肢"), (5, "記入例・備考")]:
        cell_style(ws.cell(5, col), value=label, bold=True, size=10,
                   fg=C_TITLE_FG, bg="2E75B6", h_align="center")

    r = 6  # 現在行

    # ══════════════════════════════
    # ① 会社基本情報
    # ══════════════════════════════
    section_header(ws, r, "① 会社基本情報", 0); r += 1

    items_1 = [
        ("1", "■ 会社名（正式名称）", "", "", True),
        ("2", "■ 代表者名・役職", "", "例）代表取締役　山田太郎", True),
        ("3", "■ 本社所在地", "", "例）東京都渋谷区○○1-2-3", True),
        ("4", "■ 業種・事業内容", "", "例）小売業（衣料品販売）", True),
        ("5", "■ 従業員数（常時）", "", "例）正社員15名・パート8名", True),
        ("6", "就業規則の適用範囲", "", "例）正社員・契約社員・パート（別規則あり）", False),
        ("7", "労働組合の有無", "", "あり / なし", False),
        ("8", "既存の就業規則の有無", "", "あり（改定）/ なし（新規作成）", False),
    ]
    for no, label, val, note, req in items_1:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ② 労働時間・休憩
    # ══════════════════════════════
    section_header(ws, r, "② 労働時間・休憩", 1); r += 1

    items_2 = [
        ("9",  "■ 所定労働時間（1日）", "", "例）8時間（始業9:00 終業18:00）", True),
        ("10", "■ 始業・終業時刻", "", "例）始業 9:00 / 終業 18:00", True),
        ("11", "■ 休憩時間・時刻", "", "例）12:00〜13:00（60分）", True),
        ("12", "変形労働時間制", "", "なし / 1ヶ月単位 / 1年単位 / フレックス", False),
        ("13", "フレックスタイム制", "", "あり（コアタイム○時〜○時）/ なし", False),
        ("14", "テレワーク・在宅勤務", "", "あり / なし / 一部あり", False),
        ("15", "みなし労働時間制", "", "なし / 事業場外みなし / 裁量労働", False),
        ("16", "時間外労働の上限（目安）", "", "例）月45時間以内（36協定に基づく）", False),
        ("17", "深夜・休日労働の有無", "", "あり / なし / 業務による", False),
    ]
    for no, label, val, note, req in items_2:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ③ 休日・休暇
    # ══════════════════════════════
    section_header(ws, r, "③ 休日・休暇", 2); r += 1

    items_3 = [
        ("18", "■ 週の所定休日数", "", "例）週2日（土・日）", True),
        ("19", "■ 休日の曜日・設定方法", "", "例）土日祝 / シフト制（4週8休）", True),
        ("20", "■ 年間休日数（目安）", "", "例）120日", True),
        ("21", "国民の祝日の扱い", "", "休日 / 所定労働日（振替あり）", False),
        ("22", "年次有給休暇の付与方法", "", "法定通り / 一斉付与（○月）/ 入社時即日付与", False),
        ("23", "有給の時間単位取得", "", "あり / なし", False),
        ("24", "特別休暇①　慶弔休暇", "", "例）結婚5日・忌引3日 など", False),
        ("25", "特別休暇②　その他", "", "例）子の看護休暇・介護休暇・リフレッシュ休暇", False),
        ("26", "育児・介護休業", "", "法定通り / 法定以上（内容をご記入ください）", False),
        ("27", "産前産後休業", "", "法定通り / 法定以上", False),
    ]
    for no, label, val, note, req in items_3:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ④ 賃金
    # ══════════════════════════════
    section_header(ws, r, "④ 賃金", 3); r += 1

    items_4 = [
        ("28", "■ 賃金の締切日", "", "例）毎月末日 / 毎月20日", True),
        ("29", "■ 賃金の支払日", "", "例）翌月25日 / 当月末日", True),
        ("30", "■ 賃金の支払方法", "", "銀行振込 / 現金", True),
        ("31", "基本給の種類", "", "月給制 / 日給月給制 / 時給制", False),
        ("32", "昇給", "", "あり（年1回○月）/ なし / 業績による", False),
        ("33", "賞与（ボーナス）", "", "あり（年○回：○月・○月）/ なし / 業績による", False),
        ("34", "諸手当①　役職手当", "", "あり（金額・条件をご記入ください）/ なし", False),
        ("35", "諸手当②　通勤手当", "", "例）実費支給（上限○万円）/ 定額○円", False),
        ("36", "諸手当③　家族手当", "", "あり（条件をご記入ください）/ なし", False),
        ("37", "諸手当④　住宅手当", "", "あり（条件をご記入ください）/ なし", False),
        ("38", "諸手当⑤　その他", "", "例）資格手当・皆勤手当・食事手当 など", False),
        ("39", "残業代の計算方法", "", "法定通り（時間外：1.25倍・深夜：1.25倍・休日：1.35倍）", False),
        ("40", "固定残業代（みなし残業）", "", "あり（○時間分・○円）/ なし", False),
    ]
    for no, label, val, note, req in items_4:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ⑤ 採用・試用期間
    # ══════════════════════════════
    section_header(ws, r, "⑤ 採用・試用期間", 4); r += 1

    items_5 = [
        ("41", "試用期間の有無・期間", "", "あり（○ヶ月）/ なし", False),
        ("42", "試用期間中の労働条件", "", "本採用と同じ / 賃金○%減 など", False),
        ("43", "採用時の提出書類", "", "例）履歴書・住民票・資格証明書 など", False),
        ("44", "身元保証人の要否", "", "必要 / 不要", False),
    ]
    for no, label, val, note, req in items_5:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ⑥ 服務規律
    # ══════════════════════════════
    section_header(ws, r, "⑥ 服務規律", 5); r += 1

    items_6 = [
        ("45", "副業・兼業", "", "禁止 / 届出制（許可制）/ 自由", False),
        ("46", "ハラスメント防止方針", "", "法定通り記載 / 独自の取り組みあり（詳細をご記入ください）", False),
        ("47", "SNS・情報管理のルール", "", "ガイドラインあり / 就業規則内に記載", False),
        ("48", "服装・身だしなみ規定", "", "あり（詳細をご記入ください）/ なし", False),
        ("49", "持ち物・私物PCの取り扱い", "", "禁止 / 届出制 / 自由", False),
        ("50", "競業避止義務（退職後）", "", "あり（期間・範囲をご記入ください）/ なし", False),
        ("51", "秘密保持義務", "", "あり / なし", False),
    ]
    for no, label, val, note, req in items_6:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ⑦ 懲戒
    # ══════════════════════════════
    section_header(ws, r, "⑦ 懲戒", 6); r += 1

    items_7 = [
        ("52", "懲戒の種類", "", "例）訓告・減給・出勤停止・降格・諭旨解雇・懲戒解雇", False),
        ("53", "減給の制裁の上限", "", "法定通り（1回：平均賃金の半額・総額：月給の1/10）", False),
        ("54", "懲戒事由（特記事項）", "", "例）SNS投稿・情報漏洩・無断欠勤○日以上 など", False, 30),
    ]
    for item in items_7:
        no, label, val, note, req = item[0], item[1], item[2], item[3], item[4]
        h = item[5] if len(item) > 5 else 18
        item_row(ws, r, no, label, val, note, required=req, height=h); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ⑧ 退職・解雇
    # ══════════════════════════════
    section_header(ws, r, "⑧ 退職・解雇", 7); r += 1

    items_8 = [
        ("55", "■ 自己都合退職の申出期限", "", "例）退職希望日の1ヶ月前 / 2週間前", True),
        ("56", "定年制の有無・年齢", "", "あり（○歳）/ なし", False),
        ("57", "定年後の再雇用制度", "", "あり（65歳まで）/ なし", False),
        ("58", "解雇予告", "", "法定通り（30日前予告または30日分の平均賃金）", False),
        ("59", "退職金制度", "", "あり（規程別途）/ なし / 中退共加入", False),
    ]
    for no, label, val, note, req in items_8:
        item_row(ws, r, no, label, val, note, required=req); r += 1

    ws.row_dimensions[r].height = 5; r += 1

    # ══════════════════════════════
    # ⑨ 安全衛生・その他
    # ══════════════════════════════
    section_header(ws, r, "⑨ 安全衛生・その他", 0); r += 1

    items_9 = [
        ("60", "健康診断の実施", "", "法定通り（年1回）/ 追加項目あり", False),
        ("61", "メンタルヘルス対策", "", "ストレスチェック実施 / 相談窓口設置 など", False),
        ("62", "育児・介護支援の独自措置", "", "法定以上の支援がある場合はご記入ください", False),
        ("63", "社宅・寮の有無", "", "あり（詳細をご記入ください）/ なし", False),
        ("64", "制服・作業服の貸与", "", "あり（クリーニング費用負担：会社 / 本人）/ なし", False),
        ("65", "その他　特記事項", "",
         "業界特有のルール・過去のトラブル事例・必ず盛り込みたい条文 など", False, 40),
    ]
    for item in items_9:
        no, label, val, note, req = item[0], item[1], item[2], item[3], item[4]
        h = item[5] if len(item) > 5 else 18
        item_row(ws, r, no, label, val, note, required=req, height=h); r += 1

    # ── フッター ──
    ws.row_dimensions[r].height = 5; r += 1
    merge_range(ws, r, 1, r, 5)
    cell_style(ws.cell(r, 1),
               value="【ご確認事項】このシートの情報をもとに就業規則ドラフトを作成します。"
                     "最終的な内容は社労士が確認・修正いたします。ご不明な点はご遠慮なくお申し付けください。",
               size=9, fg="1F4E79", bg="DEEAF1", border=False, wrap=True)
    ws.row_dimensions[r].height = 25

    # ── データ検証（ドロップダウン）──
    # 変形労働時間制
    add_dv(ws, '"なし,1ヶ月単位,1年単位,フレックス"',
           [ws.cell(r_n, 4) for r_n in range(6, r) if ws.cell(r_n, 1).value == "12"])
    # 副業・兼業
    add_dv(ws, '"禁止,届出制（許可制）,自由"',
           [ws.cell(r_n, 4) for r_n in range(6, r) if ws.cell(r_n, 1).value == "45"])

    # 印刷設定
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_area = f"A1:E{r}"


# ─────────────────────────────────────────────
# Sheet 2: 記入ガイド
# ─────────────────────────────────────────────
def build_guide_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30

    merge_range(ws, 1, 1, 1, 3)
    cell_style(ws.cell(1, 1), value="就業規則ヒアリングシート　記入ガイド",
               bold=True, size=13, fg=C_TITLE_FG, bg=C_TITLE_BG,
               h_align="center", border=False)

    from openpyxl.utils import get_column_letter
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 50

    guides = [
        ("■ 必須項目（赤背景）について", [
            ("会社名", "登記上の正式名称をご記入ください"),
            ("労働時間", "1日の所定労働時間・始業終業時刻は就業規則の核心部分です"),
            ("休日", "週の休日数・曜日を明確にしてください（シフト制の場合は「シフト制」と記入）"),
            ("賃金", "締切日・支払日は法律上の絶対的必要記載事項です"),
            ("退職申出", "退職希望日の何日前に申し出るかを決めてください（一般的には1ヶ月〜2ヶ月前）"),
        ]),
        ("■ よくあるご質問", [
            ("「未定」「要検討」でも大丈夫ですか？",
             "はい。まずはわかる範囲でご記入ください。後日ヒアリングで確認します"),
            ("パートと正社員で規則を分けた方がいいですか？",
             "従業員10名以上の場合、別規則の作成をお勧めします。ご相談ください"),
            ("既存の就業規則がある場合は？",
             "既存の規則をご持参いただければ、改定箇所を明示した形でドラフトを作成します"),
            ("副業禁止は有効ですか？",
             "一律禁止は現在リスクがあります。「届出制・許可制」が一般的です"),
            ("競業避止義務は盛り込めますか？",
             "可能ですが、範囲・期間・代償措置がないと無効になる場合があります"),
        ]),
        ("■ 就業規則作成の流れ", [
            ("Step 1", "このヒアリングシートにご記入いただきます"),
            ("Step 2", "社労士がドラフトを作成します（約1〜2週間）"),
            ("Step 3", "内容をご確認いただき、修正・調整を行います"),
            ("Step 4", "就業規則を完成させ、労働基準監督署へ届出します"),
            ("Step 5", "従業員への周知・説明をサポートします"),
        ]),
    ]

    row = 3
    for section, items in guides:
        ws.row_dimensions[row].height = 22
        merge_range(ws, row, 1, row, 3)
        cell_style(ws.cell(row, 1), value=section, bold=True, size=11,
                   fg=C_ITEM_FG, bg=C_ITEM_BG, border=False)
        row += 1
        for q, a in items:
            ws.row_dimensions[row].height = 20
            cell_style(ws.cell(row, 1), bg="F0F0F0")
            cell_style(ws.cell(row, 2), value=q, bold=True, size=10)
            cell_style(ws.cell(row, 3), value=a, size=10, wrap=True)
            row += 1
        row += 1


# ─────────────────────────────────────────────
# メイン
# ─────────────────────────────────────────────
def main():
    wb = openpyxl.Workbook()

    ws_main  = wb.active
    ws_main.title = "ヒアリングシート"
    ws_guide = wb.create_sheet("記入ガイド")

    build_hearing_sheet(ws_main)
    build_guide_sheet(ws_guide)

    path = "/home/user/umb/就業規則作成_ヒアリングシート.xlsx"
    wb.save(path)
    print(f"✅ 作成完了: {path}")


if __name__ == "__main__":
    main()
