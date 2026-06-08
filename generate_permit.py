from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ページ余白設定
section = doc.sections[0]
section.page_width  = Cm(21)
section.page_height = Cm(29.7)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'),  kwargs.get(edge, 'single'))
        tag.set(qn('w:sz'),   '4')
        tag.set(qn('w:space'),'0')
        tag.set(qn('w:color'),'000000')
        tcBorders.append(tag)
    tcPr.append(tcBorders)

def cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def para_font(para, bold=False, size=10.5, color=None, align=None):
    if align:
        para.alignment = align
    for run in para.runs:
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = '游明朝'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
        if color:
            run.font.color.rgb = RGBColor(*bytes.fromhex(color))

def add_header_row(table, texts, bg='1F497D', fg='FFFFFF'):
    row = table.rows[0]
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        cell_bg(cell, bg)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(*bytes.fromhex(fg))
        run.font.name = '游明朝'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')

# ===== タイトル =====
title = doc.add_paragraph('社有車使用許可証')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.runs[0]
run.bold = True
run.font.size = Pt(18)
run.font.name = '游明朝'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

doc.add_paragraph()

# 許可番号・発行日
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = info.add_run('許可番号：＿＿＿＿＿＿　　発行日：　　　年　　月　　日')
r.font.size = Pt(10)
r.font.name = '游明朝'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')

doc.add_paragraph()

# ===== セクション見出し =====
def section_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(f'■ {text}')
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = '游明朝'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

# ===== 汎用2列テーブル =====
def two_col_table(rows_data, col_widths=(4, 11)):
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.style = 'Table Grid'
    for i, (label, value) in enumerate(rows_data):
        row = table.rows[i]
        row.height = Cm(0.8)
        # ラベル列
        lc = row.cells[0]
        lc.text = label
        cell_bg(lc, 'DCE6F1')
        set_cell_border(lc)
        lc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        lr = lc.paragraphs[0].runs[0]
        lr.bold = True
        lr.font.size = Pt(10)
        lr.font.name = '游明朝'
        lr._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
        # 値列
        vc = row.cells[1]
        vc.text = value
        set_cell_border(vc)
        vc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        if vc.paragraphs[0].runs:
            vr = vc.paragraphs[0].runs[0]
            vr.font.size = Pt(10)
            vr.font.name = '游明朝'
            vr._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
    # 列幅
    for row in table.rows:
        row.cells[0].width = Cm(col_widths[0])
        row.cells[1].width = Cm(col_widths[1])
    return table

# ===== 申請者情報 =====
section_title('申請者情報')
two_col_table([
    ('所属部署',               '　'),
    ('氏名',                   '　'),
    ('社員番号',               '　'),
    ('連絡先（内線・携帯）',   '　'),
    ('運転免許証番号',         '　'),
    ('免許証有効期限',         '　　年　　月　　日'),
])

doc.add_paragraph()

# ===== 使用車両情報 =====
section_title('使用車両情報')
two_col_table([
    ('車両番号（ナンバープレート）', '　'),
    ('車種・型式',                   '　'),
    ('車両管理番号',                 '　'),
])

doc.add_paragraph()

# ===== 使用内容 =====
section_title('使用内容')
two_col_table([
    ('使用目的',           '　'),
    ('使用期間',           '　　年　　月　　日（　）　～　　年　　月　　日（　）'),
    ('使用時間',           '　　時　　分　～　　時　　分'),
    ('行先（目的地）',     '　'),
    ('経由地',             '　'),
    ('同乗者',             '　'),
    ('走行距離（予定）',   '約　　　km'),
])

doc.add_paragraph()

# ===== 誓約事項 =====
section_title('誓約事項')
pledges = [
    '道路交通法その他関連法規を遵守し、安全運転に努めます。',
    '飲酒・薬物服用後の運転は絶対に行いません。',
    '使用前後に車両点検を行います。',
    '事故・違反が発生した場合は、直ちに会社に報告します。',
    '業務目的以外には使用しません。',
    '車内の整理整頓を保ち、貴重品を車内に放置しません。',
    '使用後は燃料を補給し、所定の場所に返却します。',
]
for i, text in enumerate(pledges, 1):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = '游明朝'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')

doc.add_paragraph()

# ===== 申請・承認欄 =====
section_title('申請・承認欄')
approval_table = doc.add_table(rows=3, cols=5)
approval_table.style = 'Table Grid'
headers = ['　', '申請者', '直属上長', '部門長', '総務担当']
for i, h in enumerate(headers):
    cell = approval_table.rows[0].cells[i]
    cell.text = h
    cell_bg(cell, 'DCE6F1')
    set_cell_border(cell)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cell.paragraphs[0].runs:
        r = cell.paragraphs[0].runs[0]
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = '游明朝'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')

row_labels = ['氏名・印', '確認日']
for ri, label in enumerate(row_labels, 1):
    row = approval_table.rows[ri]
    row.height = Cm(1.5)
    for ci in range(5):
        cell = row.cells[ci]
        set_cell_border(cell)
        if ci == 0:
            cell.text = label
            cell_bg(cell, 'DCE6F1')
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cell.paragraphs[0].runs:
                r = cell.paragraphs[0].runs[0]
                r.bold = True
                r.font.size = Pt(10)
                r.font.name = '游明朝'
                r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')

doc.add_paragraph()

# ===== 使用実績 =====
section_title('使用実績（返却時記入）')
two_col_table([
    ('返却日時',   '　　年　　月　　日　　時　　分'),
    ('実走行距離', '　　　km'),
    ('燃料補給',   '□ 補給済み　□ 不要'),
    ('車両の異常', '□ 異常なし　□ 異常あり（内容：　　　　　　　　　　　）'),
    ('備考',       '　'),
])

doc.add_paragraph()

# ===== 注意事項 =====
note = doc.add_paragraph()
note.paragraph_format.left_indent  = Cm(0.5)
note.paragraph_format.right_indent = Cm(0.5)
r = note.add_run('【注意事項】　本許可証は使用期間中、必ず携帯してください。許可なく第三者に車両を運転させることは禁止します。不正使用が判明した場合は懲戒処分の対象となります。')
r.font.size = Pt(9)
r.font.name = '游明朝'
r._element.rPr.rFonts.set(qn('w:eastAsia'), '游明朝')
r.font.color.rgb = RGBColor(0x80, 0x00, 0x00)

out = '/home/user/umb/社有車使用許可証.docx'
doc.save(out)
print(f'Saved: {out}')
