from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# スタイル設定
style = doc.styles['Normal']
style.font.name = 'メイリオ'
style.font.size = Pt(10.5)

# タイトル
title = doc.add_heading('AIを始めよう ― AI活用ヒント・手順まとめ', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

doc.add_paragraph('出典：https://office-roumu1.com/ai-hajimeru（2026年4月時点の情報）\n')

# ─── はじめに ───
doc.add_heading('はじめに', 1)
doc.add_paragraph(
    'このサイトは、AIエンジニアではない文系の著者が試行錯誤しながら見つけた「AI活用のヒントと手順」をまとめた場所です。'
    '「全部読まなくていい、気になるところから読んでください」というスタンスで書かれています。'
    '著者自身が「ITスキルがなくてもやりたいことを形にできるなら使いこなしたい」という動機でAIを試してきた実体験が元になっています。'
)

# ─── これからの時代 ───
doc.add_heading('① これからの時代のこと', 1)
doc.add_paragraph(
    '少子高齢化による労働人口の減少が続く日本では、AIの活躍する場面が拡大していきます。'
    '経済産業省の試算では、2040年に向けて以下の大きな変化が見込まれています。'
)
tbl_data = [
    ('事務職の余剰', '約440万人'),
    ('AI人材の不足', '約340万人'),
    ('文系人材の余剰', '約80万人'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light List Accent 1'
hdr = table.rows[0].cells
hdr[0].text = '項目'
hdr[1].text = '規模感'
for item, val in tbl_data:
    row = table.add_row().cells
    row[0].text = item
    row[1].text = val
doc.add_paragraph('')
doc.add_paragraph(
    '定型的な事務作業はAIに置き換わっていく一方、AI人材は大幅に不足する見通しです。'
    'AIを「敵」ではなく「味方」にすることが、事務職・文系の人間が時代を生き抜くカギだと著者は述べています。'
)

# ─── 主なAIサービス ───
doc.add_heading('② 主なAIサービスと使い分け', 1)
doc.add_paragraph('「どのAIを使えばいいか」迷ったときの目安です。')

services = [
    ('ChatGPT（OpenAI）', '万能型。会話・文章作成・アイデア出し。無料版あり。まず始めるならここ。'),
    ('Claude（Anthropic）', '長文読解・精密作業が得意。Claude CodeはPCファイルを直接操作して作業を代行。'),
    ('Gemini（Google）', 'Gmail・Googleドキュメント・スプレッドシートとの連携が強み。画像・動画理解も高性能。'),
    ('Copilot（Microsoft）', 'Word・Excel・PowerPoint・Outlook内でAIが使える。Microsoft 365利用者に最適。'),
    ('NotebookLM（Google）', '自分のアップした資料だけを根拠に回答。マニュアルや規程集の質問に最適。'),
    ('Genspark', '複数AIモデルを同時に使える。検索結果を1ページのレポートにまとめる「Sparkpage」機能あり。'),
    ('Gamma', 'プレゼン・スライド作成に特化。テーマを伝えるだけでデザイン済みスライドを自動生成。'),
    ('Grok（xAI）', 'X（旧Twitter）と連携。リアルタイムのトレンドに強い。'),
]
tbl2 = doc.add_table(rows=1, cols=2)
tbl2.style = 'Light List Accent 1'
h = tbl2.rows[0].cells
h[0].text = 'サービス名'
h[1].text = '特徴・得意分野'
for name, desc in services:
    r = tbl2.add_row().cells
    r[0].text = name
    r[1].text = desc
doc.add_paragraph('')

doc.add_heading('場面別おすすめ', 2)
use_cases = [
    'ちょっとした質問・雑談 → ChatGPT（無料版でOK）',
    '手元の資料について質問 → NotebookLM',
    'Googleツール内で使いたい → Gemini',
    'Excel・Wordの中で使いたい → Copilot',
    '本格的な文書作成・ファイル操作・ツール開発 → Claude Code',
    'ウェブ検索を1ページのレポートにまとめたい → Genspark',
    'プレゼン・スライドをサッと作りたい → Gamma',
]
for u in use_cases:
    p = doc.add_paragraph(u, style='List Bullet')

# ─── Claudeについて ───
doc.add_heading('③ Claude / Anthropic社について', 1)
doc.add_paragraph(
    'Anthropic社はOpenAIの元メンバーが「より安全なAI」を目指して設立したアメリカの企業です（本社：サンフランシスコ）。'
)
facts = [
    '企業価値：約3,800億ドル（2026年2月時点）≒ 約57兆円',
    'Google出資額：約30億ドル',
    '年率売上：約300億ドル規模（2026年4月時点）',
]
for f in facts:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('Claudeのサービス構成', 2)
services2 = [
    ('Claude チャット', 'ブラウザで使える会話型AI。ちょっとした質問・相談に。'),
    ('Claude Code', 'PCのファイルを直接操作できるAI。CLI版とデスクトップアプリ版がある。'),
    ('Claude Cowork', 'Claude CodeをGUI操作に寄せた非エンジニア向けのサービス。'),
]
tbl3 = doc.add_table(rows=1, cols=2)
tbl3.style = 'Light List Accent 1'
h3 = tbl3.rows[0].cells
h3[0].text = 'サービス'
h3[1].text = '説明'
for name, desc in services2:
    r = tbl3.add_row().cells
    r[0].text = name
    r[1].text = desc
doc.add_paragraph('')

doc.add_heading('モデルと料金の目安', 2)
models = [
    ('Haiku', '軽量・高速。簡単な質問向き。'),
    ('Sonnet', 'バランス型。最初はこれで十分。'),
    ('Opus', '最高性能。複雑な作業に。'),
]
tbl4 = doc.add_table(rows=1, cols=2)
tbl4.style = 'Light List Accent 1'
h4 = tbl4.rows[0].cells
h4[0].text = 'モデル'
h4[1].text = '特徴'
for name, desc in models:
    r = tbl4.add_row().cells
    r[0].text = name
    r[1].text = desc
doc.add_paragraph('')
doc.add_paragraph('料金：Pro プラン 月額$20（約3,000円）／ Max プラン 月額$100〜（約15,000円〜）')
doc.add_paragraph('＊料金の差は「AIに処理してもらえる文字量（トークン）の上限」の違いです。')

# ─── Claude Codeが別格な理由 ───
doc.add_heading('④ Claude Codeが「別格」に感じる理由', 1)
doc.add_paragraph(
    'ChatGPTなど他のAIチャットと比較したときのClaude Codeの優位点として、著者は以下を挙げています。'
)
reasons = [
    '長編小説2〜3冊分の文章を数秒で読む処理速度',
    '数十〜数百のファイルを同時に把握しながら整合チェックができる広さ',
    'ハルシネーション（もっともらしい嘘）率が大幅に低下した正確さ',
    'チャットで答えを返すだけでなく「PCのファイルを操作して作業そのものをやってくれる」点',
]
for r in reasons:
    doc.add_paragraph(r, style='List Bullet')

doc.add_heading('実際に数日で作ったもの（著者実績）', 2)
made = [
    '行政のWord見本 → 入力しやすいExcelへ変換',
    '雇用契約を管理するExcelシート',
    'ホームページへの就業規則自動診断機能の搭載',
    'ホームページのデザイン更新・セミナー資料の作成',
    'Chromeブラウザ操作の自動化',
    'パソコンのバックアップ自動化',
]
for m in made:
    doc.add_paragraph(m, style='List Bullet')

# ─── 他AIとの違い ───
doc.add_heading('⑤ GPTs・GemとClaude Codeの違い', 1)
doc.add_paragraph(
    '結論：GPTs（ChatGPT）・Gem（Gemini）は「クラウド上」で動き、Claude Codeは「あなたのPC上」で動きます。'
    'そのためClaude Codeは、PC内のファイルに直接アクセスして仕事を進められます。'
    'これが実務で大きな差になります。'
)

# ─── セットアップ ───
doc.add_heading('⑥ セットアップ・基本設定', 1)

doc.add_heading('推奨ブラウザ', 2)
doc.add_paragraph('Google Chrome を推奨（Claude in Chrome などの拡張機能がChrome専用のため）。')

doc.add_heading('Gitのインストールについて', 2)
doc.add_paragraph(
    'WindowsでClaude Codeを使うにはGitのインストールが公式必須です。'
    'ただし、ユーザー自身がGitの「履歴を残す・前に戻す」機能を能動的に使う必要はなく、'
    'インストールさえしておけばClaude Codeが裏で必要に応じて使ってくれます。'
)

doc.add_heading('CLAUDE.md（基本設定ファイル）', 2)
doc.add_paragraph('Claude Codeを「自分仕様」にするための設定ファイルが3つあります。')
files = [
    ('CLAUDE.md', 'AIに「私はこういう人です」と教えるファイル。毎回の説明が不要になる。'),
    ('about-me.md', 'プロフィール情報。'),
    ('brand-voice.md', '文体・トーンの指定。'),
]
tbl5 = doc.add_table(rows=1, cols=2)
tbl5.style = 'Light List Accent 1'
h5 = tbl5.rows[0].cells
h5[0].text = 'ファイル名'
h5[1].text = '用途'
for name, desc in files:
    r = tbl5.add_row().cells
    r[0].text = name
    r[1].text = desc
doc.add_paragraph('')
doc.add_paragraph(
    '設定後、Claude Codeに「私のプロフィールを教えて」と聞いてCLAUDE.mdの内容が返ってくれば設定完了です。'
)

# ─── 活用例 ───
doc.add_heading('⑦ 具体的な活用例', 1)

doc.add_heading('ホームページ作成', 2)
doc.add_paragraph(
    'Claude Codeに「どんなサイトを作りたいか」を日本語で伝えるだけで、'
    'HTMLでデザインしてくれます。著者はブルー・ベージュ・ゴールド・グリーン・ピンクと色違いの見本サイトを作成済みです。'
)

doc.add_heading('パンフレット・チラシ作成', 2)
steps = [
    'Claude Codeに「A4二つ折りの事務所案内パンフレットを作って」と伝える',
    '事務所名・住所・サービス内容・連絡先などを箇条書きで渡す',
    'ClaudeがHTMLでデザインしてくれるのでブラウザで開いてPDF印刷',
    '「もう少しやわらかい雰囲気で」など何度でも修正OK',
]
for i, s in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {s}')

doc.add_heading('セミナー資料作成', 2)
doc.add_paragraph(
    'テーマを伝えるだけでたたき台が1〜2時間で完成。'
    'ただし法律の正確性が命のため、プロンプトに「一次情報源を明示させる」「法改正日を確認させる」縛りを入れることが重要。'
    '著者はPowerPoint一式・構成案＋原稿のテキスト・Word配布資料を使い分けています。'
)

doc.add_heading('Webアプリ作成', 2)
apps = [
    '就業規則の自動診断（質問に答えるだけで改善ポイントがわかる）',
    'お客様向けのかんたんアンケートフォーム（回答を自動集計）',
    '料金シミュレーター（従業員数を入れると顧問料の目安がわかる）',
]
for a in apps:
    doc.add_paragraph(a, style='List Bullet')
doc.add_paragraph('完成したらNetlify Dropにアップしてそのままシェアすることも可能です。')

# ─── Claudeへの上手な頼み方 ───
doc.add_heading('⑧ Claudeに上手にお願いするコツ', 1)
doc.add_paragraph(
    'Anthropic公式の勉強会で紹介された「指示の出し方」を押さえると、返ってくる答えが大きく変わります。'
    'ポイントは基本5要素・応用10要素・XMLタグでの整理。'
    '開発者直伝のコツとして「まず質問から始める」「計画してから動く」「CLAUDE.mdに自分ルールを書く」なども有効です。'
)

# ─── セッション管理 ───
doc.add_heading('⑨ セッション管理のコツ', 1)
doc.add_paragraph(
    'セッションが長くなったら「引継ぎ書を作って」と頼み、新しいセッションで「この引継ぎを読んで続きを」と伝えるとスムーズです。'
)
doc.add_paragraph('疲れてきたサイン：「圧縮します」表示 ／ 返答が遅い ／ 同じ間違いの繰り返し → 新しいセッションに切り替えましょう。')

# ─── セキュリティ ───
doc.add_heading('⑩ セキュリティで気をつけること', 1)
sec = [
    '学習させない設定 ― Settingsで学習をオフにする',
    '機密情報・マイナンバー・パスワード・クレジットカード番号はAIに入力しない',
    '出力は必ず人間がチェック（AIも間違えます）',
    '自動チャージの設定を確認しておく',
    'コネクタは大手公式のものだけ接続する',
    '固有名詞は仮名に置き換えるなどの工夫を',
]
for s in sec:
    doc.add_paragraph(s, style='List Bullet')

# ─── 組織での活用 ───
doc.add_heading('⑪ 組織でAI活用を始めるときの心がまえ', 1)
doc.add_paragraph(
    '反対する人がいるのは健全。電子申請・パソコン普及のときも同じでした。'
    '大切なのは以下の3つです。'
)
for item in ['"なぜ取り入れたいか"を共有する', '小さく気負わず試す', '遊び心を持つ']:
    doc.add_paragraph(item, style='List Bullet')

# ─── 用語集 ───
doc.add_heading('⑫ 用語集', 1)
terms = [
    ('セッション', 'Claude Codeとの1回の会話のこと'),
    ('トークン', 'AIが処理する文字量の単位'),
    ('プロンプト', 'AIへの指示や質問のこと'),
    ('CLAUDE.md', 'AIに自分のことを覚えてもらうための設定ファイル'),
    ('スキル（SKILL.md）', 'よく使う作業の「指示書」'),
    ('コネクタ', '外部サービスとClaudeをつなげる仕組み'),
    ('Git', 'ファイルの変更履歴を記録する仕組み'),
    ('コンテキスト', 'AIが今把握している情報の量'),
    ('MCP', 'コネクタの技術名（「コネクタのことね」でOK）'),
    ('ハルシネーション', 'AIがもっともらしい嘘をつく現象'),
]
tbl6 = doc.add_table(rows=1, cols=2)
tbl6.style = 'Light List Accent 1'
h6 = tbl6.rows[0].cells
h6[0].text = '用語'
h6[1].text = '意味'
for term, meaning in terms:
    r = tbl6.add_row().cells
    r[0].text = term
    r[1].text = meaning
doc.add_paragraph('')

# ─── 免責事項 ───
doc.add_heading('免責事項', 1)
doc.add_paragraph(
    'このページの内容は著者の個人的な体験・工夫であり、公式ガイドや法的助言ではありません。'
    'AI業界の変化は早く情報が古くなる場合があります。'
    '大事な判断の際は最新の公式情報を確認するか専門家に相談してください。'
    'このページの情報を参考にした結果生じた損害・責任については、著者は責任を負いません。'
)

# 保存
output_path = '/home/user/umb/AI活用ヒント手順まとめ.docx'
doc.save(output_path)
print(f'保存完了: {output_path}')
