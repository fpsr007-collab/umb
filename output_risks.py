import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "就業規則未整備リスク一覧"

# ── データ定義 ────────────────────────────────────────────────
data = [
    # (No, 項目名, 説明, 想定リスク, is_new, category)
    (1,  "懲戒処分が根拠薄弱・法的に無効となるリスク",
         "懲戒事由・処分の種類・手続きが明確でないため、無断欠勤・業務命令違反・情報持ち出し・ハラスメントなどへの対処の根拠が弱くなります。裁判所・労働審判の判例（フジ興産事件・最高裁2003年）でも、就業規則への明記が懲戒処分の有効要件とされています。",
         "問題社員対応が後手に回る・懲戒無効・バックペイ・訴訟費用リスク",
         False, "法的リスク"),
    (2,  "解雇・本採用拒否・雇止めの法的根拠が弱くなる",
         "能力不足・勤務態度不良・協調性欠如・試用期間中の適格性判断などについて会社判断が後出しに見えやすく、解雇理由・手続き・予告期間のルールが書面で存在しないため、「不当解雇」を主張された際に会社側の根拠が著しく弱くなります。",
         "解雇無効・金銭解決・不当解雇主張・労働審判・訴訟リスク",
         False, "法的リスク"),
    (3,  "助成金・制度導入で不利になる",
         "厚生労働省の多くの雇用関係助成金では就業規則の整備・届出が申請要件に含まれており、整備していないだけで申請資格を失います。規程整備・周知・制度運用が求められる助成金では申請準備も遅れます。",
         "助成金受給機会の逸失・申請準備遅延リスク",
         False, "経営リスク"),
    (4,  "休職・復職・自然退職の処理が曖昧になる",
         "私傷病で働けなくなった場合の休職期間・診断書・復職判断・休職期間満了時の取り扱いが定まらず、長期化・紛争化しやすくなります。",
         "長期欠勤・復職トラブル",
         False, "労務管理リスク"),
    (5,  "労働条件の変更が難しくなる",
         "賃金・手当・休日・勤務時間・賞与・退職金などの変更時に、基準や手続きが曖昧になり、不利益変更トラブルにつながりやすくなります。",
         "不利益変更トラブル",
         False, "法的リスク"),
    (6,  "「言った・言わない」が増える",
         "欠勤連絡・有給申請・退職届・貸与物返還・服務規律などが口頭ルールになり、従業員から「聞いていない」と言われやすくなります。",
         "社長判断への不満・不公平感",
         False, "労務管理リスク"),
    (7,  "固定残業代（みなし残業）の無効リスク",
         "固定残業代を給与に含める場合、計算根拠・対象時間数・超過時の追加払いルールが就業規則または労働条件通知書に明記されていることが有効要件のひとつです。",
         "未払い残業代・遡及請求リスク",
         False, "法的リスク"),
    (8,  "変形労働時間制・フレックスタイム制を導入できない",
         "シフト制や繁閑対応の柔軟な勤務体系を導入するには、就業規則への記載が労働基準法上の要件となっています（第32条の2・32条の4等）。",
         "制度無効・残業代発生・人件費増加リスク",
         False, "経営リスク"),
    (9,  "採用力・人材定着率の低下",
         "転職経験のある求職者は就業規則の開示を求めることがあります。開示できない場合、職場環境が不透明と判断され、優秀な人材の採用に影響します。",
         "採用競争力低下・早期離職リスク",
         False, "経営リスク"),
    (10, "パワハラ社員への対応が困難になる",
         "ハラスメントの定義・禁止行為・相談窓口・調査手続き・懲戒処分が就業規則に明記されていないと、行為者への注意・処分の根拠が曖昧になります。2022年4月からすべての事業主にパワハラ防止措置が義務付けられており（労働施策総合推進法）、対応が不十分な場合は会社が使用者責任を問われます。",
         "使用者責任・損害賠償請求・被害者離職・職場環境悪化リスク",
         True, "法的リスク"),
    (11, "能力不足社員の対応・解雇が後手に回る",
         "求められる業務水準・評価基準・指導プロセス・改善期間が就業規則や関連規程に定められていないと、能力不足を理由とする解雇の正当性を立証することが困難になります。裁判所は能力不足解雇に厳格であり、指導・教育・配置転換を行った証拠がなければ解雇無効とされるケースが多くあります。",
         "解雇無効・問題社員の長期在籍・職場モラル低下リスク",
         True, "労務管理リスク"),
    (12, "未払い残業代の遡及請求リスクが高まる",
         "労働時間の管理方法・時間外労働の認定基準・残業申請ルールが就業規則に定められていないと、実態と賃金台帳が乖離しやすくなります。2020年の労基法改正により未払い賃金の請求権消滅時効が3年に延長されており、退職者・在職者を問わず一括請求されるリスクが増大しています。",
         "3年分の未払い残業代・遅延損害金・複数社員による集団請求リスク",
         True, "法的リスク"),
    (13, "長期欠勤社員の出口戦略が取れない",
         "欠勤が続く社員に対して、休職発令の基準・休職期間の上限・診断書提出義務・復職判断の手順・期間満了時の退職扱いが就業規則に定められていないと、対応が「待ちの状態」のまま際限なく長期化します。解雇に踏み切っても、根拠規定がなければ無効とされるリスクがあります。",
         "給与負担の長期化・業務穴埋めコスト・解雇無効・自然退職扱い無効リスク",
         True, "労務管理リスク"),
]

# ── スタイル定義 ──────────────────────────────────────────────
thin  = Side(style="thin",   color="BBBBBB")
thick = Side(style="medium", color="1F3864")
BORDER     = Border(left=thin, right=thin, top=thin, bottom=thin)
BORDER_TOP = Border(left=thin, right=thin, top=thick, bottom=thin)

FILL_HEADER   = PatternFill("solid", fgColor="1F3864")   # 濃紺
FILL_NEW      = PatternFill("solid", fgColor="E2EFDA")   # 薄緑（新規追加）
FILL_OLD_A    = PatternFill("solid", fgColor="F2F6FC")   # 薄青（既存・奇数）
FILL_OLD_B    = PatternFill("solid", fgColor="FFFFFF")   # 白（既存・偶数）
FILL_RISK_H   = PatternFill("solid", fgColor="FFE0E0")   # 薄赤（法的リスク）
FILL_RISK_M   = PatternFill("solid", fgColor="FFF2CC")   # 薄黄（経営リスク）
FILL_RISK_L   = PatternFill("solid", fgColor="EAF0FF")   # 薄青紫（労務管理）
FILL_CAT_H    = PatternFill("solid", fgColor="C00000")   # 深紅（法的）
FILL_CAT_M    = PatternFill("solid", fgColor="ED7D31")   # オレンジ（経営）
FILL_CAT_L    = PatternFill("solid", fgColor="2E75B6")   # 青（労務）

FONT_H   = Font(name="Meiryo UI", bold=True, color="FFFFFF", size=10)
FONT_B   = Font(name="Meiryo UI", size=10, color="404040")
FONT_BD  = Font(name="Meiryo UI", bold=True, size=10, color="1F3864")
FONT_NEW = Font(name="Meiryo UI", bold=True, size=10, color="375623")
FONT_CAT = Font(name="Meiryo UI", bold=True, size=9,  color="FFFFFF")

WRAP   = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

# ── 列幅・行高さ ───────────────────────────────────────────────
col_widths = {"A": 5, "B": 12, "C": 32, "D": 48, "E": 36, "F": 5}
for col, w in col_widths.items():
    ws.column_dimensions[col].width = w

# ── タイトル行 ────────────────────────────────────────────────
ws.merge_cells("A1:F1")
title_cell = ws["A1"]
title_cell.value = "就業規則未整備リスク一覧　（全13項目）"
title_cell.font      = Font(name="Meiryo UI", bold=True, size=14, color="FFFFFF")
title_cell.fill      = FILL_HEADER
title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 36

# ── 凡例行 ────────────────────────────────────────────────────
ws.merge_cells("A2:B2")
ws["A2"].value = "【凡例】"
ws["A2"].font  = Font(name="Meiryo UI", bold=True, size=9, color="404040")
ws["A2"].alignment = CENTER

legend = [
    ("C2", FILL_RISK_H, "法的リスク"),
    ("D2", FILL_RISK_M, "経営リスク"),
    ("E2", FILL_RISK_L, "労務管理リスク"),
    ("F2", FILL_NEW,    "新規追加項目"),
]
for cell_addr, fill, label in legend:
    c = ws[cell_addr]
    c.value     = label
    c.fill      = fill
    c.font      = Font(name="Meiryo UI", bold=True, size=9, color="404040")
    c.alignment = CENTER
    c.border    = BORDER
ws.row_dimensions[2].height = 18

# ── ヘッダー行 ────────────────────────────────────────────────
headers = ["No", "カテゴリ", "項目名", "説　明", "想定リスク", ""]
ws.row_dimensions[3].height = 24
for ci, h in enumerate(headers, start=1):
    c = ws.cell(row=3, column=ci, value=h)
    c.font      = FONT_H
    c.fill      = FILL_HEADER
    c.alignment = CENTER
    c.border    = BORDER

# カテゴリ色マップ
CAT_FILL = {
    "法的リスク":   (FILL_RISK_H, FILL_CAT_H),
    "経営リスク":   (FILL_RISK_M, FILL_CAT_M),
    "労務管理リスク": (FILL_RISK_L, FILL_CAT_L),
}

# ── データ行 ──────────────────────────────────────────────────
for idx, (no, title, desc, risk, is_new, cat) in enumerate(data):
    row = idx + 4

    row_fill, cat_fill = CAT_FILL[cat]
    base_fill = FILL_NEW if is_new else row_fill

    # 行高さ：説明文の文字数に応じて動的設定
    lines = max(len(desc) // 28 + 1, len(risk) // 18 + 1, 3)
    ws.row_dimensions[row].height = max(lines * 15, 48)

    # A: No
    c = ws.cell(row=row, column=1, value=no)
    c.font = Font(name="Meiryo UI", bold=True, size=11,
                  color="375623" if is_new else "1F3864")
    c.fill      = base_fill
    c.alignment = CENTER
    c.border    = BORDER

    # B: カテゴリ（縦書き風・背景色でカテゴリ識別）
    c = ws.cell(row=row, column=2, value=cat)
    c.font      = Font(name="Meiryo UI", bold=True, size=9, color="FFFFFF")
    c.fill      = cat_fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center",
                            text_rotation=90)
    c.border    = BORDER

    # C: 項目名
    prefix = "★NEW　" if is_new else ""
    c = ws.cell(row=row, column=3, value=prefix + title)
    c.font      = FONT_NEW if is_new else FONT_BD
    c.fill      = base_fill
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    c.border    = BORDER

    # D: 説明
    c = ws.cell(row=row, column=4, value=desc)
    c.font      = FONT_B
    c.fill      = base_fill
    c.alignment = WRAP
    c.border    = BORDER

    # E: 想定リスク
    c = ws.cell(row=row, column=5, value=risk)
    c.font      = Font(name="Meiryo UI", bold=True, size=10,
                       color="C00000" if "訴訟" in risk or "無効" in risk or "請求" in risk
                       else "404040")
    c.fill      = base_fill
    c.alignment = WRAP
    c.border    = BORDER

    # F: 新規マーカー列
    c = ws.cell(row=row, column=6, value="NEW" if is_new else "")
    c.font      = Font(name="Meiryo UI", bold=True, size=8, color="375623")
    c.fill      = FILL_NEW if is_new else base_fill
    c.alignment = CENTER
    c.border    = BORDER

# ── ウィンドウ枠固定（ヘッダー固定） ────────────────────────
ws.freeze_panes = "A4"

# ── オートフィルター ─────────────────────────────────────────
ws.auto_filter.ref = "A3:F3"

# ── 保存 ─────────────────────────────────────────────────────
output = "/home/user/umb/就業規則未整備リスク一覧.xlsx"
wb.save(output)
print(f"Saved: {output}")
